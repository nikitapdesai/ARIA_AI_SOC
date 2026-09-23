"""
File location in project: generate_report_excel.py
    python generate_report_excel.py
"""

import json
import os

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUTS_DIR = "outputs"
OUT_PATH = "outputs/ARIA_Model_Reports.xlsx"

MODEL_DIRS = [
    ("MLP", "mlp"),
    ("MLP + SMOTE", "mlp_smote"),
    ("Random Forest", "random_forest"),
    ("Random Forest + SMOTE", "random_forest_smote"),
    ("Ensemble - MLP alone", "ensemble/mlp_on_final_eval"),
    ("Ensemble - RF alone", "ensemble/rf_on_final_eval"),
    ("Ensemble - Stacked", "ensemble/stacked_ensemble"),
]

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=13)


def load_classification_report(folder):
    path = f"{OUTPUTS_DIR}/{folder}/classification_report.json"
    if not os.path.exists(path):
        return None
    with open(path) as f:
        report = json.load(f)

    rows = []
    for label, metrics in report.items():
        if label == "accuracy":
            continue
        rows.append({
            "class": label,
            "precision": metrics["precision"],
            "recall": metrics["recall"],
            "f1-score": metrics["f1-score"],
            "support": metrics["support"],
        })
    df = pd.DataFrame(rows)
    accuracy = report.get("accuracy")
    return df, accuracy


def load_confusion_matrix(folder):
    path = f"{OUTPUTS_DIR}/{folder}/confusion_matrix.csv"
    if not os.path.exists(path):
        return None
    return pd.read_csv(path, index_col=0)


def load_feature_importance(folder):
    path = f"{OUTPUTS_DIR}/{folder}/feature_importance.csv"
    if not os.path.exists(path):
        return None
    return pd.read_csv(path)


def load_training_history(folder):
    path = f"{OUTPUTS_DIR}/{folder}/training_history.json"
    if not os.path.exists(path):
        return None
    with open(path) as f:
        history = json.load(f)
    return pd.DataFrame(history)


def style_header_row(ws, row_idx, n_cols):
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def autofit_columns(ws, df, start_col=1):
    for i, col in enumerate(df.columns):
        col_letter = get_column_letter(start_col + i)
        max_len = max(len(str(col)), df[col].astype(str).map(len).max() if len(df) else 0)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 40)


def write_df(writer, sheet_name, df, title, start_row=0, index=False):
    ws_book = writer.book
    if sheet_name not in writer.sheets:
        ws_book.create_sheet(sheet_name)
    ws = writer.sheets[sheet_name] if sheet_name in writer.sheets else ws_book[sheet_name]

    ws.cell(row=start_row + 1, column=1, value=title).font = TITLE_FONT

    df.to_excel(writer, sheet_name=sheet_name, startrow=start_row + 2, index=index)
    header_row = start_row + 3
    n_cols = len(df.columns) + (1 if index else 0)
    style_header_row(ws, header_row, n_cols)
    autofit_columns(ws, df.reset_index() if index else df, start_col=1)

    return start_row + 2 + len(df) + 4


def build_summary_rows():
    rows = []
    for sheet_name, folder in MODEL_DIRS:
        result = load_classification_report(folder)
        if result is None:
            continue
        report_df, accuracy = result
        macro = report_df.loc[report_df["class"] == "macro avg"].iloc[0] if "macro avg" in report_df["class"].values else None
        weighted = report_df.loc[report_df["class"] == "weighted avg"].iloc[0] if "weighted avg" in report_df["class"].values else None
        rows.append({
            "Model": sheet_name,
            "Accuracy": round(accuracy, 4) if accuracy is not None else None,
            "Macro Precision": round(macro["precision"], 4) if macro is not None else None,
            "Macro Recall": round(macro["recall"], 4) if macro is not None else None,
            "Macro F1": round(macro["f1-score"], 4) if macro is not None else None,
            "Weighted F1": round(weighted["f1-score"], 4) if weighted is not None else None,
        })
    return pd.DataFrame(rows)


def main():
    print("Building ARIA model report workbook ...")

    with pd.ExcelWriter(OUT_PATH, engine="openpyxl") as writer:
        summary_df = build_summary_rows()
        summary_df.to_excel(writer, sheet_name="Summary", startrow=2, index=False)
        ws = writer.sheets["Summary"]
        ws.cell(row=1, column=1, value="ARIA -- Model Evaluation Summary").font = TITLE_FONT
        style_header_row(ws, 3, len(summary_df.columns))
        autofit_columns(ws, summary_df)
        ws.freeze_panes = "A4"
        print("  Summary sheet built.")

        for sheet_name, folder in MODEL_DIRS:
            result = load_classification_report(folder)
            if result is None:
                print(f"  Skipping {sheet_name} (no classification_report.json found in {folder}/)")
                continue
            report_df, accuracy = result

            safe_name = sheet_name[:31]
            writer.book.create_sheet(safe_name)
            next_row = write_df(
                writer, safe_name, report_df,
                f"{sheet_name} -- Classification Report (accuracy = {accuracy:.4f})" if accuracy else f"{sheet_name} -- Classification Report",
            )

            cm_df = load_confusion_matrix(folder)
            if cm_df is not None:
                next_row = write_df(writer, safe_name, cm_df, "Confusion Matrix (rows = true, columns = predicted)", start_row=next_row, index=True)

            fi_df = load_feature_importance(folder)
            if fi_df is not None:
                next_row = write_df(writer, safe_name, fi_df.head(20), "Top 20 Feature Importances", start_row=next_row)

            hist_df = load_training_history(folder)
            if hist_df is not None:
                next_row = write_df(writer, safe_name, hist_df, "Training History (per epoch)", start_row=next_row)

            writer.sheets[safe_name].freeze_panes = "A3"
            print(f"  {sheet_name} sheet built.")

    print(f"\nSaved: {OUT_PATH}")


if __name__ == "__main__":
    main()
